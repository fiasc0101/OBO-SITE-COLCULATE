import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COLORS = {"ok": "E3F3E1", "review": "FFF3D1", "none": "FFFFFF", "decided": "EEEEEE", "acc": "E6F0FA"}

def build_excel(rows, with_price=False):
    wb = Workbook(); ws = wb.active; ws.title = "Подбор"
    head = ["№", "Наименование (проект)", "Артикул (проект)", "Ед.", "Кол-во",
            "Наименование (OBO)", "Тип (буквенный артикул)", "Артикул (цифровой)", "Ед.", "Кол-во", "Примечание"]
    if with_price: head += ["Цена", "Сумма"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="1A1A1A")
    thin = Side(style="thin", color="999999")

    def add(vals, color, price=None, qty=None):
        if with_price:
            vals += [price, (price * qty) if price is not None and qty is not None else None]
        ws.append(vals)
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=color); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(wrap_text=True, vertical="top")

    for i, r in enumerate(rows, 1):
        s, m = r["src"], r.get("match")
        if m:
            right = [m["name"], m["letter_article"], m["num_article"], m["unit"], m["qty"], r.get("how", "")]
        elif r.get("decision"):
            right = [r["decision"], "", "", "", "", ""]
        else:
            right = ["", "", "", "", "", "Не подобрано"]
        add([s.get("pos") or i, s["name"], s.get("article"), s.get("unit"), s["qty"]] + right,
            COLORS.get(r["status"], "FFFFFF"), m.get("price") if m else None, m["qty"] if m else None)
        for a in r.get("accessories", []):
            add(["", "  ↳ крепёж / соединитель", "", "", "", a["name"], a["letter_article"], a["num_article"],
                 a["unit"], a["qty"], a.get("note") or "Закладывается автоматически"],
                COLORS["acc"], a.get("price"), a["qty"])
    for col, wd in zip("ABCDEFGHIJKLM", [6, 40, 18, 7, 10, 48, 20, 16, 7, 10, 26, 10, 12]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
